"""Build matching Word and Markdown reports from the completed analysis."""
from __future__ import annotations

import json
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
RES = ROOT / "results"
OUT = ROOT / "report" / "Assignment3_Optimization_Report.docx"
TITLE = "Optimizing a Southeastern U.S. Distribution Network"


def money(value, decimals=0):
    return f"${value:,.{decimals}f}"


def figure_analyses(summary, comparison, capacity, simulation, sensitivity):
    """Compose a 50-word interpretation of each figure using the saved run."""
    opt, data = summary["milp"], summary["instance"]
    methods = comparison.set_index("method")
    savings = 1 - opt["network_cost"] / methods.loc["Baseline (6 metro hubs)", "network_cost"]
    rated = capacity.set_index("dc").capacity_shadow_price
    nominal, hedged = simulation.iloc[0], simulation.iloc[1]
    raised_truck_cost = sensitivity[(sensitivity.parameter == "Truck cost per mile")
                                     & (sensitivity.level == 1.15)].iloc[0]
    analyses = {
        "fig_travel_time.png": (
            "Travel times generally increase with estimated distance, but variation "
            f"remains between county pairs. Only {data['allowed_arc_share']:.1%} of candidate routes "
            f"satisfy the modeled {summary['run']['parameters']['sla_minutes'] / 60:g} hour median "
            "threshold, restricting feasible assignments. The visible spread supports using route "
            "specific time estimates. This restriction makes geographic coverage important "
            "when selecting distribution centers and allocating demand."),
        "fig_network_map.png": (
            f"The selected network uses {opt['n_dcs']} centers serving modeled demand at "
            f"{money(opt['network_cost'])} weekly. Allocation balances inbound transportation, "
            f"outbound distance, fixed costs, and travel times. Its {savings:.1%} cost "
            "advantage over the six-hub benchmark accompanies compliance with the median "
            "time screen. The counties provide a shortlist for evaluating "
            "properties, transport access, and coverage."),
        "fig_utilization.png": (
            f"Utilization ranges from {opt['util_min']:.0%} to {opt['util_max']:.0%} across the "
            "selected centers, indicating uneven workloads within the permitted band. Orange and "
            "Polk reach the upper limit, while Leon operates at the minimum. Additional capacity "
            "has different value across sites. Capacity planning should prioritize "
            "constrained centers while accounting for each location's minimum throughput."),
        "fig_model_comparison.png": (
            "Exact optimization provides the lowest feasible cost for this candidate set. "
            "Simulated annealing and repaired gradient search match that reference, while "
            "greedy methods remain close. LP rounding opens additional centers and increases "
            "cost. The MILP combines the best cost with a short solve time, making "
            "it the strongest reference method."),
        "fig_sa_convergence.png": (
            "The best simulated annealing score approaches the MILP reference as location "
            "changes improve the network. Temporary increases in current cost reflect accepted "
            "exploratory moves. The final feasible network matches the reference cost. The "
            "trajectory illustrates how accepting temporary deterioration helps search "
            "different location combinations before retaining the best feasible allocation."),
        "fig_gradient.png": (
            "Continuous site trajectories show how gradient updates explore potential locations "
            "before snapping to candidate counties. The changing Lagrangian reflects temperature "
            "and capacity price updates during search. Discrete repair closes the cost gap "
            "after snapping. This result supports combining geographic exploration "
            "with feasible network evaluation and local adjustments to warehouse selections."),
        "fig_sensitivity.png": (
            "Increasing truck costs by 15% changes the selected network from "
            f"{opt['n_dcs']} to {raised_truck_cost.n_dcs:.0f} centers. Cost and location choices "
            "respond differently across the parameter sweeps. Capacity changes also alter "
            "minimum throughput requirements, explaining why expansion need not reduce "
            "cost. Planners should consider transportation "
            "economics and throughput policies when evaluating network changes."),
        "fig_shadow_prices.png": (
            "Orange and Polk offer rated capacity savings of "
            f"{money(rated['Orange, FL'], 2)} and {money(rated['Polk, FL'], 2)} per additional "
            "weekly pallet, respectively. Leon has a negative capacity value because expansion "
            "also raises minimum throughput. Demand marginal costs vary geographically. These "
            "fixed-network marginal values identify where small capacity adjustments "
            "offer savings and prioritize expansion options for analysis."),
        "fig_monte_carlo.png": (
            f"Across {summary['run']['simulation_scenarios']} simulated demand scenarios, the nominal "
            f"design fills {nominal.mean_fill_rate:.3%} of demand on average. The design for "
            "higher demand records no sampled shortages, with mean penalized cost increasing "
            f"by {hedged.mean_cost / nominal.mean_cost - 1:.2%}. This comparison supports "
            "evaluating capacity flexibility. Planning should weigh the facility "
            "commitments against lower shortage exposure under the tested conditions."),
    }
    for filename, analysis in analyses.items():
        words = len(analysis.split())
        if words != 50:
            raise ValueError(f"{filename}: analysis must have 50 words, found {words}")
    return analyses


class Report:
    def __init__(self, analyses):
        self.doc = Document()
        self.markdown = []
        self.analyses = analyses
        self.n_table = self.n_figure = 0
        normal = self.doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(7)
        normal.paragraph_format.line_spacing = 1.15
        for section in self.doc.sections:
            section.left_margin = section.right_margin = Inches(1)
            section.top_margin = section.bottom_margin = Inches(0.8)

    def h(self, text, level=1):
        self.doc.add_heading(text, level=level)
        self.markdown.append(f"{'#' * (level + 1)} {text}\n")

    def p(self, text, style=None):
        paragraph = self.doc.add_paragraph(style=style)
        for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
            run = paragraph.add_run(part.strip("*") if part.startswith("*") else part)
            run.bold = part.startswith("**")
            run.italic = part.startswith("*") and not part.startswith("**")
        self.markdown.append(text + "\n")
        return paragraph

    def eq(self, text):
        paragraph = self.p(text)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.name = "Cambria Math"

    def table(self, frame, caption, font=8):
        self.n_table += 1
        self.p(f"**Table {self.n_table}. {caption}**")
        table = self.doc.add_table(rows=1, cols=len(frame.columns))
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, column in zip(table.rows[0].cells, frame.columns):
            cell.text = str(column)
        header_repeat = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(header_repeat)
        for values in frame.itertuples(index=False, name=None):
            for cell, value in zip(table.add_row().cells, values):
                cell.text = str(value)
        for index, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    for run in paragraph.runs:
                        run.font.size = Pt(font)
                        run.bold = index == 0
        def clean(value):
            return str(value).replace("|", "/").replace("\n", " ")
        self.markdown.append("| " + " | ".join(map(clean, frame.columns)) + " |")
        self.markdown.append("|" + "|".join("---" for _ in frame.columns) + "|")
        self.markdown.extend("| " + " | ".join(map(clean, values)) + " |"
                             for values in frame.itertuples(index=False, name=None))
        self.markdown.append("")
        self.doc.add_paragraph()

    def figure(self, filename, caption, width=6.2):
        analysis = self.analyses[filename]
        self.n_figure += 1
        self.doc.add_picture(str(RES / "figures" / filename), width=Inches(width))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
        self.markdown.append(f"![{caption}](../results/figures/{filename})\n")
        self.p(f"*Figure {self.n_figure}. {caption}*").paragraph_format.keep_with_next = True
        self.p(analysis).paragraph_format.keep_together = True

    def reference(self, text):
        paragraph = self.p(text)
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.5)

    def save(self):
        properties = self.doc.core_properties
        properties.title = TITLE
        properties.subject = "Assignment 3: Optimization Techniques"
        properties.author = properties.last_modified_by = "Albert Kabore"
        properties.comments = properties.keywords = properties.category = ""
        output = OUT
        try:
            self.doc.save(output)
        except PermissionError:
            # Windows locks documents open in Word. Keep the completed revision
            # available without closing the user's document or losing the changes.
            output = OUT.with_stem(OUT.stem + "_Updated")
            self.doc.save(output)
        markdown = "\n".join(self.markdown)
        OUT.with_suffix(".md").write_text(markdown, encoding="utf-8")
        if output != OUT:
            output.with_suffix(".md").write_text(markdown, encoding="utf-8")
        return output


def build():
    summary = json.loads((RES / "summary.json").read_text(encoding="utf-8"))
    run, data, opt, brute = (summary[k] for k in ("run", "instance", "milp", "brute_force"))
    comparison = pd.read_csv(RES / "model_comparison.csv")
    iterations = pd.read_csv(RES / "iterations.csv")
    tornado = pd.read_csv(RES / "tornado.csv")
    sensitivity = pd.read_csv(RES / "sensitivity.csv")
    capacity = pd.read_csv(RES / "capacity_shadow_prices.csv")
    simulation = pd.read_csv(RES / "monte_carlo_summary.csv")
    audit = pd.read_csv(RES / "feasibility_audit.csv")
    inventory = pd.read_csv(RES / "inventory_policy.csv")
    cleaning = pd.read_csv(RES / "cleaning_audit.csv")
    lookup = comparison.set_index("method")
    baseline = lookup.loc["Baseline (6 metro hubs)"]
    gradient = lookup.loc["Gradient (snap + repair)"]
    annealing = lookup.loc["Simulated annealing"]
    savings = 1 - opt["network_cost"] / baseline.network_cost
    nominal, hedged = simulation.iloc[0], simulation.iloc[1]
    p = run["parameters"]
    R = Report(figure_analyses(summary, comparison, capacity, simulation, sensitivity))
    R.doc.add_heading(TITLE, 0)
    R.markdown.append(f"# {TITLE}\n")
    R.p("Assignment 3: Optimization Techniques")
    R.p("Albert Kabore")

    R.h("Abstract")
    R.p(f"This exercise selects distribution-center counties and weekly shipment volumes for a "
        f"hypothetical network serving {data['customers']} counties in six southeastern states. "
        f"Census population and warehousing statistics establish the regional pattern, and BTS "
        f"county-pair truck times provide a travel-time proxy. A mixed-integer linear program (MILP), "
        f"an LP relaxation, greedy search, simulated annealing, and a gradient-based location method "
        f"are compared. The MILP selects {opt['n_dcs']} of {data['candidates']} candidate counties "
        f"at {money(opt['network_cost'])} per week, including {money(opt['transport_cost'])} in "
        f"transportation cost. This is {savings:.1%} below the six-hub benchmark. "
        f"All allocated volume uses arcs whose modeled median time is at most eight hours. "
        f"The LP provides a "
        f"lower bound, while heuristic and gradient results are assessed against the MILP and "
        f"independent feasibility checks. Sensitivity analysis and {run['simulation_scenarios']} "
        f"simulated demand scenarios examine cost, capacity, and service trade-offs.")

    R.h("Introduction")
    R.p("Warehouse placement determines how shipments move from supply points to customers. "
        "Additional warehouses can shorten outbound trips, but they also add operating costs and "
        "may increase inbound travel. The planning problem is to choose locations and flows "
        "together while meeting demand, capacity, and service requirements.")
    R.p("The study covers Alabama, Florida, Georgia, North Carolina, South Carolina, and Tennessee. "
        "The 60 most populous counties form demand zones. A manually selected list of 22 counties "
        "defines the available facility choices. Chatham County, Georgia, represents a single "
        "supply point associated with Savannah.")
    R.p("The objectives are to formulate a facility-location problem, minimize transportation and "
        "facility operating costs, control utilization across centers, evaluate delivery-time "
        "restrictions, and compare linear, integer, heuristic, and gradient-based methods. "
        "The six assignment steps are implemented as data preparation, stated operational "
        "assumptions, model development, validation, model iteration, and reporting.")

    R.h("Methodology")
    R.h("Data collection and preprocessing", 2)
    R.p("The source inputs are the 2019–2023 ACS total-population table B01003, 2022 County "
        "Business Patterns (CBP) for NAICS 493, 2020 county population mean centers, the 2023 "
        "county Gazetteer, and the experimental 2024 BTS county-pair truck-time dataset "
        "(U.S. Census Bureau, n.d., 2023b, 2024a, 2024b; Bureau of Transportation Statistics, 2026). "
        "Five-digit county FIPS codes join the files. The CBP extract keeps the all-size, "
        "all-legal-form rows. Missing source rows remain unknown; the selected candidate "
        "counties must have reported establishments and employment.")
    verified_text = ("All compared values matched the live source records." if run["source_values_verified"]
                     else "The current raw files do not have a complete matching source-verification record.")
    R.p(f"Source verification compared 534 Census population rows, 522 CBP rows across employment "
        f"size classes, all 1,300 downloaded BTS rows, and the two geography files. {verified_text} "
        "The verification results and file hashes are recorded in results/source_verification.json.")
    description = pd.DataFrame([
        ["ACS B01003, 2019–2023", "Population; persons", "Keep the largest 60 counties in six states; demand = population × 0.0015 pallets/week (assumed)."],
        ["CBP 2022, NAICS 493", "Establishments; employees", "County warehousing activity; 22 candidate counties chosen for this exercise. An assumed 3,000-employee threshold assigns capacity tiers."],
        ["2020 population centers", "Latitude/longitude; degrees", "Population mean centers represent the county nodes."],
        ["2023 Gazetteer", "Land area; square miles", "Density = population / land area; assumed facility-cost scaling uses density capped at 3,000 persons/sq mi."],
        ["BTS truck times, 2024", "County-pair p25/p50/p75; minutes", f"1,300 source pairs; {data['bts_rows_adjusted']} adjusted with a centroid-distance floor; {data['assumed_local_pairs']} local pairs assumed 30 minutes; {data['imputed_nonlocal_pairs']} nonlocal imputations."],
        ["ATRI 2024 cost study", "Truck operating cost; $/mile", "Historical 2023 benchmark $2.270/mile; converted to pallet-mile costs using assumed loads and empty mileage."],
    ], columns=["Source", "Variables and units", "Preprocessing / use"])
    R.table(description, "Data description", font=8)
    R.h("Data cleaning and validation", 2)
    R.p("Data preparation is executed before optimization and recorded in "
        "results/cleaning_audit.csv. County codes are trimmed and normalized to five "
        "digits, required numeric fields are parsed, and duplicate or missing keys "
        "are rejected. Population and business counts must be finite, nonnegative "
        "integers; flagged CBP totals require explicit treatment. Geographic joins "
        "are checked one-to-one, with valid coordinates and positive land area. "
        "Missing geography stops processing instead of silently removing counties.")
    R.table(cleaning[["stage", "input_rows", "output_rows", "affected_rows", "type"]].rename(
        columns={"stage": "Stage", "input_rows": "Input", "output_rows": "Output",
                 "affected_rows": "Affected", "type": "Operation type"}),
        "Executed cleaning audit; full check descriptions are in cleaning_audit.csv", font=8)
    qc = cleaning.set_index("check")
    R.p(f"Selecting CBP county totals retains "
        f"{qc.loc['Keep all-size / all-legal-form NAICS 493 totals', 'output_rows']:,} of "
        f"{qc.loc['Keep all-size / all-legal-form NAICS 493 totals', 'input_rows']:,} rows. "
        f"The remaining rows contain employment-size breakdowns. "
        f"{qc.loc['Counties without a CBP total', 'affected_rows']:,} study counties have no "
        "CBP total in this extract; those values remain unknown. Every selected candidate "
        "has reported employment and establishments. No duplicates or long BTS elapsed "
        "times were removed. BTS year, nonnegative finite quartiles, and quartile order "
        "are checked. Local-time assignments and the distance-based floor are flagged "
        "in the processed route file, alongside the original source times.")
    R.p(f"The selected counties contain {data['population']:,} residents in the source population "
        f"estimates, producing {data['demand_pallets_week']:,.1f} assumed pallets per week. "
        "Population supplies the demand weights, while CBP employment assigns the warehouse "
        "capacity tiers. The source years and the operational assumptions are listed in "
        "the data and input tables.")
    R.p("BTS measures stopped and moving elapsed time from the last truck observation in an "
        "origin county to the first in a destination county. Pairs with fewer than 100 "
        "observed movements are excluded, and movement counts are binned. The published "
        "quartiles summarize annual county-pair travel times "
        "(Bureau of Transportation Statistics, n.d.).")
    R.p(f"Estimated road distance is {p['circuity']:.2f} times the great-circle distance between "
        "county centers. For this exercise, the p50 and p75 time proxies are floored at "
        "estimated miles / 55 mph × 60, with p75 constrained to be no less than p50. "
        "This adjustment aligns the time proxy with the estimated county-center distance. "
        "Raw times remain in the processed route file. "
        f"All {data['assumed_local_pairs']} missing pairs are same-county movements; each receives "
        "an assumed 30 minutes. The generic nonlocal imputation routine is unused in this instance.")
    R.figure("fig_travel_time.png", "Modeled county-pair times and estimated distances; the horizontal line is the eight-hour exercise threshold.", width=5.5)

    R.h("Operational assumptions and prior knowledge", 2)
    assumptions = pd.DataFrame([
        ["Demand", f"{p['pallets_per_1000_pop']} pallets per 1,000 residents per week", "Assumed population-to-demand conversion"],
        ["Facility tiers", "Large: 12,000; small: 7,000 pallets/week", "Assumed; county employment only assigns the tier"],
        ["Facility fixed costs", "$60,000 / $38,000 per week", "Assumed; density factor ranges from 0.85 to 1.15"],
        ["Truck cost", "$2.270 per vehicle-mile", "ATRI historical 2023 operating-cost benchmark"],
        ["Loads / empty mileage", "26 pallet positions; outbound 70% full; inbound 0.15 empty miles/loaded mile", "Assumed average truck loading and mileage factors"],
        ["Handling", "$6 per pallet", "Assumed constant per pallet handled"],
        ["Utilization / time", "35%–90%; modeled p50 ≤ 480 minutes", "Assumed throughput band and arc-screening rule"],
        ["Inventory policy", "One-week ordering cycle; 0.5-week lead time; z = 1.645; $4/pallet-week", "Approximate inventory indicators; holding cost excluded from network objective"],
        ["Demand uncertainty", "County shock CV 25%; shared regional shock CV 8%", "Assumed mean-one lognormal shocks; shortages priced at $500/pallet"],
    ], columns=["Input", "Value", "Basis / use"])
    R.table(assumptions, "Economic and operational inputs", font=8)
    R.p("The truck-cost benchmark is the 2023 average in ATRI's 2024 update (Leslie & Murray, "
        "2024; American Transportation Research Institute, 2024). Outbound cost is 2.270 / (26 × 0.70) = "
        f"{money(data['outbound_rate'], 4)} per pallet-mile. Inbound cost is 2.270 × 1.15 / 26 = "
        f"{money(data['inbound_rate'], 4)} per pallet-mile. The 15% allowance means empty miles "
        "per loaded mile. Costs use estimated distance, while the time restriction uses the "
        "county-pair time proxy.")

    R.h("Mathematical formulation", 2)
    R.p("Let i index candidate counties and j index demand counties. Binary yᵢ opens a "
        "distribution center. Continuous xᵢⱼ is the fraction of county j's demand assigned to i. "
        "Demand dⱼ is measured in pallets/week, Kᵢ is rated weekly throughput capacity, "
        "fᵢ is weekly fixed cost, and cᵢⱼ is inbound plus handling plus outbound dollars/pallet.")
    for equation in [
        "minimize  Σᵢ fᵢyᵢ + ΣᵢΣⱼ cᵢⱼdⱼxᵢⱼ",
        "Σᵢ xᵢⱼ = 1  for each j",
        "0.35 Kᵢyᵢ ≤ Σⱼ dⱼxᵢⱼ ≤ 0.90 Kᵢyᵢ  for each i",
        "xᵢⱼ ≤ yᵢ;  0 ≤ xᵢⱼ ≤ 1;  yᵢ ∈ {0,1}",
        "xᵢⱼ = 0 whenever modeled tᵢⱼ⁵⁰ > 480 minutes",
    ]:
        R.eq(equation)
    R.p("The model has 22 binary location variables and 1,320 continuous assignment variables. "
        "It permits split sourcing and fractional pallets in weekly aggregate flows. Supply "
        "is assumed sufficient; inbound volume equals each center's outbound throughput. "
        "The objective includes fixed facilities, inbound transport, handling, and outbound "
        "transport. Inventory holding is calculated separately from the weekly network "
        "objective.")
    R.p("The throughput band prevents opening a center with arbitrarily little assigned volume. "
        "Under the common one-week ordering-cycle policy, average cycle stock is "
        "Iᵢ = 0.5 Σⱼdⱼxᵢⱼ. Thus the throughput band also bounds normalized cycle stock: "
        "0.175 Kᵢyᵢ ≤ Iᵢ ≤ 0.45 Kᵢyᵢ. Inventory levels are calculated after routing. "
        "Safety stock uses z√(lead time × weekly variance), with both county-specific and "
        "shared regional variation included. The normal approximation uses z = 1.645 "
        "and the nominal routing allocations. The separate demand simulation evaluates "
        "capacity and shipment allocation for each demand realization.")

    R.h("Algorithms and comparison design", 2)
    R.p("The reference MILP is solved with HiGHS through SciPy, with a relative optimality-gap "
        "tolerance of 10⁻⁶. Replacing binary y with 0 ≤ y ≤ 1 gives the LP lower bound. "
        "The LP-based rounding heuristic opens every location with y > 10⁻⁶, then solves the "
        "fixed-network transportation LP (SciPy, n.d.-a, n.d.-b). Every resulting network "
        "is audited for demand, throughput, capacity, and time constraints.")
    R.p("Greedy ADD starts with no open centers and accepts the best cost-reducing opening. "
        "Greedy DROP starts with all centers open and accepts the best closing. Simulated "
        f"annealing uses {run['sa_iterations']:,} flip/swap proposals, seed 42, and a geometric "
        "temperature schedule from 0.02 to 0.00001. An uphill move is accepted with "
        "exp[−Δ / (temperature × current cost)] (Kirkpatrick et al., 1983). Each network is "
        "scored by an LP in pallet-flow units. Search scores permit unmet demand and "
        "minimum-throughput shortfall at $500/pallet. Costs are compared with the MILP only "
        "after an independent audit confirms feasibility.")
    R.p("The gradient method fits K continuous sites in an approximate planar coordinate system. "
        "Its smooth transport surrogate includes inbound/outbound distance, a soft distance "
        "penalty for the time proxy, and capacity prices. Let gⱼₖ be these per-pallet terms "
        "including capacity price μₖ, and C be the common surrogate capacity. The Lagrangian is:")
    R.eq("L(z, μ) = Σⱼ dⱼ [−τ log Σₖ exp(−gⱼₖ/τ)] − C Σₖ μₖ")
    R.p(f"Site positions use Adam updates (Kingma & Ba, 2015); nonnegative capacity prices use "
        f"projected dual ascent. K ranges from 5 to 12 with {run['gradient_restarts']} restarts "
        "and 1,500 iterations per restart. Temperature falls from 20 to 0.2. The distance "
        f"penalty begins around {summary['gradient_sla_radius_miles']:.0f} estimated road miles, "
        "based on a coarse empirical time-screening rule. A linear assignment maps the fitted "
        "points to distinct candidate counties. A local repair step opens or closes sites to "
        "reduce unmet demand and throughput deficits. Feasible candidates are ranked by exact "
        "network cost. The continuous stage approximates geography, capacities, and travel "
        "times; the discrete evaluation adds site-specific fixed costs and enforces the "
        "network constraints.")
    R.p("The methods address different parts of the problem. The MILP supplies a certified "
        "reference for this finite candidate list; the LP quantifies a lower bound; the "
        "heuristics demonstrate discrete search; and the gradient method tests a smooth "
        "location approximation followed by discrete evaluation. They are compared on "
        "feasibility, cost, selected centers, and elapsed computation time.")

    R.h("Validation, iteration, and sensitivity design", 2)
    R.p("Independent checks cover demand equality, maximum capacity, minimum throughput, the "
        "time-proxy screen, shipments from open centers only, and nonnegative flows. "
        f"For a {brute['candidates']}-candidate reduced instance, every nonempty open set is "
        "evaluated with hard demand and throughput constraints and compared with the MILP. "
        "Automated tests also compare demand/capacity marginal values with numerical cost "
        "changes and analytic gradients with finite differences.")
    R.p("The iteration study adds capacity, the time screen, minimum throughput, and then "
        "single sourcing. Six one-at-a-time parameter sweeps are supplemented by ±20% shocks "
        "to seven cost and operating inputs. Shadow prices come from a fixed-network LP; "
        "they describe local changes while the open-center set is held constant.")
    R.p(f"The simulation uses {run['simulation_scenarios']} paired scenarios, seed 7. Each county's "
        "demand is multiplied by independent mean-one lognormal shocks for that county and "
        "the whole region. The combined county CV is about "
        f"{(((1 + p['demand_cv']**2) * (1 + p['common_demand_cv']**2) - 1)**0.5):.1%}. "
        "For each scenario, the same fixed-network LP reallocates flows and prices shortages. "
        "The two designs are the nominal MILP and a MILP designed for 15% higher demand.")

    R.h("Results")
    R.h("Selected network and cost", 2)
    R.p(f"The MILP selects {opt['n_dcs']} centers: **{opt['dcs']}**. The solver reports a relative "
        f"gap of {opt['mip_gap']:.2g}; solve time is {opt['runtime_s']:.3f} seconds on this run. "
        f"Weekly network cost is {money(opt['network_cost'])}, or {money(opt['cost_per_pallet'], 2)} "
        "per pallet.")
    R.table(pd.DataFrame({"Cost component": ["Facility fixed", "Inbound transportation", "Outbound transportation", "Handling", "Total network"],
                          "$/week": [money(opt[k]) for k in ["fixed", "inbound", "outbound", "handling", "network_cost"]]}),
            "Optimal weekly cost")
    R.p(f"Transportation alone costs {money(opt['transport_cost'])} per week. Average modeled "
        f"outbound distance is {opt['avg_miles']:.1f} estimated road miles. The flow-weighted "
        f"average of arc p50 time proxies is {opt['avg_time_p50_min']/60:.2f} hours. All assigned volume uses "
        f"arcs passing the p50 screen, while {opt['pct_within_sla_p75']:.1%} uses arcs whose "
        "p75 proxy also passes eight hours.")
    R.p("Map backgrounds use Census 2023 state cartographic boundaries for geographic "
        "orientation (U.S. Census Bureau, 2023a). Lines connect the selected center "
        "and demand nodes, with line widths representing assigned weekly volumes.")
    R.figure("fig_network_map.png", "Selected county centers and weekly flows. Lines show shipment assignments; the star marks the Chatham County supply proxy.", width=5.5)
    R.table(pd.DataFrame({"Center": inventory.dc, "Pallets/week": inventory.throughput.map(lambda x: f"{x:,.0f}"),
                          "Utilization": inventory.utilization.map(lambda x: f"{x:.1%}"),
                          "Cycle stock": inventory.cycle_stock.map(lambda x: f"{x:,.0f}"),
                          "Safety stock (approx.)": inventory.safety_stock_approx.map(lambda x: f"{x:,.0f}")}),
            "Throughput and inventory-policy indicators; stock units are pallets")
    R.p(f"Cycle stock totals {opt['cycle_stock']:,.0f} pallets and approximate safety stock "
        f"{opt['safety_stock']:,.0f} pallets. At the assumed holding rate, their reported "
        f"weekly holding cost is {money(opt['inventory_holding'])}; it is excluded from the "
        "network objective. Inventory quantities scale with throughput under the common coverage policy.")
    R.figure("fig_utilization.png", "Throughput utilization relative to the assumed 35%–90% band.", width=6.0)

    R.h("LP, heuristic, and gradient comparisons", 2)
    relaxation = summary["lp_relaxation"]
    R.p(f"The continuous LP relaxation is a lower bound of {money(relaxation['objective'])} "
        f"per week. The difference (MILP − LP) / LP is {relaxation['integrality_gap_pct']:.3f}%. "
        f"There are {relaxation['fractional_y']} fractional facility variables.")
    R.table(pd.DataFrame({"Method": comparison.method,
                          "Centers": comparison.n_dcs,
                          "Cost incl. search penalties": comparison.objective_incl_penalties.map(money),
                          "Gap*": comparison.gap_vs_milp_pct.map(lambda x: "—" if pd.isna(x) else f"{x:+.2f}%"),
                          "Feasible*": comparison.feasible_base_model.map({True: "Yes", False: "No"}),
                          "Time (s)": comparison.runtime_s.map(lambda x: f"{x:.3f}")}),
            "Method comparison. *Feasibility and gaps use all base-model constraints.", font=7.5)
    R.p("The six-hub benchmark uses Jefferson, Duval, Orange, Fulton, Mecklenburg, and "
        "Davidson counties and is routed without the time screen. Feasible-solution "
        "rankings include only methods that pass all base-model constraints. The cost "
        "column includes any search penalties. Runtimes are elapsed seconds for the "
        "recorded run on the platform listed in the appendix. "
        "The two gradient rows share the same multi-restart search time.")
    R.figure("fig_model_comparison.png", "Cost gaps and runtimes for solutions that pass every base-model constraint.")
    R.figure("fig_sa_convergence.png", "Simulated-annealing search scores and the MILP reference; scores include any unmet-demand or throughput penalties.", width=5.5)
    R.figure("fig_gradient.png", "Gradient trajectories and the smoothed Lagrangian as capacity prices and temperature change.")

    R.h("Validation and model iteration", 2)
    n_pass = int(audit.drop(columns="method").all(axis=1).sum())
    R.p(f"The independent audit passes for {n_pass} of {len(audit)} optimized solutions. "
        f"The reduced enumeration examines {brute['subsets']:,} nonempty subsets, of which "
        f"{brute['feasible_subsets']} are feasible. Its best cost is "
        f"{money(brute['enumeration_cost'], 2)}, compared with {money(brute['milp_cost'], 2)} "
        f"from the reduced MILP; the relative difference is {brute['rel_diff']:.2g}. "
        "Exhaustive enumeration checks the reduced solution, while the full-instance "
        "MILP reports its termination status and optimality gap.")
    R.table(audit.rename(columns=lambda x: x.replace("_", " ")).replace({True: "Pass", False: "Fail"}),
            "Independent constraint checks", font=7.5)
    R.table(pd.DataFrame({"Version": iterations.version, "Centers": iterations.n_dcs,
                          "$/week": iterations.network_cost.map(money),
                          "p50 proxy ≤ 8 h": iterations.pct_within_sla_p50.map(lambda x: f"{x:.1%}"),
                          "Utilization CV": iterations.util_cv.map(lambda x: f"{x:.3f}")}),
            "Constraint additions and single-sourcing variant", font=8)
    R.p("The first four versions successively add restrictions to the same split-sourcing "
        "problem. The fifth requires each county to use a single center. A constraint can "
        "increase minimum cost while changing the selected site count in either direction. "
        "The cost differences quantify the effect of each added restriction.")

    R.h("Sensitivity and simulated demand", 2)
    R.table(pd.DataFrame({"Input": tornado.parameter,
                          "Cost change at −20%": tornado.low.map(lambda x: f"{x:+.2%}"),
                          "Cost change at +20%": tornado.high.map(lambda x: f"{x:+.2%}")}),
            "Reoptimized one-at-a-time changes")
    R.figure("fig_sensitivity.png", "Reoptimized cost across parameter sweeps. Labels give open-center counts; infeasible settings are marked.")
    R.table(pd.DataFrame({"Center": capacity.dc,
                          "Upper-bound relaxation ($/pallet)": capacity.usable_capacity_shadow_price.map(lambda x: f"{x:.2f}"),
                          "Rated-capacity change ($/pallet/week)": capacity.capacity_shadow_price.map(lambda x: f"{x:.2f}")}),
            "Local LP capacity savings for a fixed open-center set", font=8)
    R.p("The first marginal-value column relaxes the usable upper throughput bound alone. "
        "The second increases rated K by one pallet/week, changing the upper bound by 0.90 "
        "and the lower bound by 0.35. These values can differ and can change when an "
        "active constraint changes. Demand marginal values are obtained with demand on the "
        "right-hand side of a pallet-flow LP; small finite-difference checks verify their units.")
    R.figure("fig_shadow_prices.png", "Fixed-network demand marginal costs and local rated-capacity savings.", width=6.0)
    R.table(pd.DataFrame({"Design": simulation.network, "Centers": simulation.n_dcs,
                          "Mean penalized cost": simulation.mean_cost.map(money),
                          "95th pct. cost": simulation.p95_cost.map(money),
                          "Mean fill": simulation.mean_fill_rate.map(lambda x: f"{x:.3%}"),
                          "P(unmet > 1 pallet)": simulation.prob_shortage.map(lambda x: f"{x:.1%}")}),
            f"Demand simulation: {run['simulation_scenarios']} paired scenarios", font=8)
    R.figure("fig_monte_carlo.png", "Simulated costs and quantities filled on arcs that pass the time-proxy screen.", width=6.0)

    R.h("Analysis and Discussion")
    R.p(f"The six-hub benchmark costs {money(baseline.network_cost)} per week, "
        f"against {money(opt['network_cost'])} for the selected design, a difference of "
        f"{money(baseline.network_cost-opt['network_cost'])} ({savings:.1%}). The comparison "
        f"network places {baseline.pct_within_sla_p50:.1%} of its flow on arcs passing the "
        "p50 screen, compared with 100% for the selected design. Jointly choosing "
        "warehouse locations and shipment assignments reduces weekly cost while "
        "meeting the median-time constraint.")
    R.p(f"Inbound transport accounts for {opt['inbound']/opt['network_cost']:.1%} of selected "
        "network cost. The supply node is co-located with the Chatham candidate, giving "
        "that center zero modeled inbound mileage. This treatment highlights the "
        "influence of the supply-point location on facility selection. A property-level "
        "evaluation can add port drayage distance and location-specific handling costs.")
    R.p(f"Simulated annealing finishes {annealing.gap_vs_milp_pct:+.2f}% from the reference "
        f"cost in this seeded run; the repaired gradient method finishes "
        f"{gradient.gap_vs_milp_pct:+.2f}% from it. The MILP supplies the optimality "
        "certificate and has the shortest runtime among these three methods. "
        "The gradient approach combines a smooth geographic surrogate with discrete "
        "repair, which improves the snapped network to the reference cost. This "
        "improvement shows why feasibility repair is valuable after continuous fitting.")
    ranked = tornado.assign(impact=tornado[["low", "high"]].abs().max(axis=1)).sort_values("impact", ascending=False)
    R.p("The largest changes in reoptimized cost under the tested ±20% shocks come from "
        + ", ".join(ranked.parameter.head(3).str.lower()) + ". Each sweep changes one "
        "parameter at a time. Changing capacity also "
        "changes the minimum-throughput requirement, so more rated capacity need not "
        "lower cost under that policy. Circuity changes cost distances while the main "
        "time-proxy matrix is held fixed. The resulting cost changes measure the "
        "effect of distance assumptions on network design.")
    R.p(f"Under the assumed demand distribution, nominal mean fill is {nominal.mean_fill_rate:.3%}, "
        f"and {nominal.prob_shortage:.1%} of scenarios have more than one unmet pallet. "
        f"The design for 15% higher demand has {hedged.n_dcs:.0f} centers and mean fill "
        f"{hedged.mean_fill_rate:.3%}, with {(hedged.mean_cost/nominal.mean_cost-1):+.2%} "
        "mean penalized cost relative to the nominal design. Flow reallocation adapts "
        "shipments to each scenario. The higher-demand design exchanges a small increase "
        "in mean cost for improved fill rates across the tested scenarios.")
    R.p("The formulation operates at county level with weekly aggregate flows. "
        "Property-level sites, delivery schedules, driver hours, loading queues, "
        "and multi-period inventory balances are useful extensions for operational "
        "planning. Correlated transport disruptions can also be added to the "
        "demand uncertainty analysis.")

    R.h("Recommendations")
    recommendations = [
        "Use the selected counties as a shortlist for further investigation. Obtain actual "
        "facility addresses, availability, lease and labor costs, dock capacity, storage "
        "capacity, and port access to compare properties within each county.",
        "Replace the population-based pallet rate with historical customer orders and "
        "observed shipment sizes. Fit demand forecasts on past periods and evaluate them "
        "on later periods before passing forecasts and uncertainty ranges into the optimizer.",
        "Estimate route-specific delivery distributions from dispatch and arrival records. "
        "Choose the service percentile and promised delivery window explicitly, including "
        "loading, unloading, and driver constraints. Use BTS county summaries as "
        "regional benchmarks alongside trip-level records.",
        "Evaluate extra capacity where fixed-network marginal savings are positive, then "
        "reoptimize with realistic expansion increments and costs. Use the demand-growth "
        "design to assess the costs and benefits of preparing for demand growth.",
        "Add a multi-period inventory model if stock balancing is a business requirement. "
        "Include opening inventory, replenishment, customer shipments, storage limits, "
        "holding cost, and shortages. The present throughput band only controls normalized "
        "cycle stock under the stated coverage policy.",
        "Connect validated demand and travel predictions to the MILP, monitor forecast "
        "error and constraint violations, and rerun the analysis as orders and operating "
        "conditions change. Test the resulting plans on historical periods and compare "
        "their costs and service outcomes with the current planning process.",
    ]
    for number, text in enumerate(recommendations, 1):
        R.p(f"{number}. {text}")

    R.h("Conclusion")
    R.p(f"For the defined historical-data scenario, the MILP selects {opt['n_dcs']} county "
        f"centers at {money(opt['network_cost'])} per week and satisfies the demand, capacity, "
        "throughput, and modeled median-time constraints. The LP supplies a lower bound; "
        "simulated annealing and the repaired gradient method match the MILP cost and pass "
        "the feasibility audit. Sensitivity analysis identifies demand, transport cost, "
        "and distance as major cost drivers. The demand-growth design improves sampled "
        "fill rates with a small increase in mean penalized cost. Together, these results "
        "support evaluating warehouse placement, capacity flexibility, and service "
        "requirements as connected planning decisions.")

    R.h("References")
    references = [
        "American Transportation Research Institute. (2024, June 25). *New ATRI research: "
        "Industry costs increased more than 6 percent during freight recession.* "
        "https://truckingresearch.org/2024/06/new-atri-research-industry-costs-increased-more-than-6-percent-during-freight-recession/",
        "Bureau of Transportation Statistics. (n.d.). *Trucking movements: BTS Freight Mobility "
        "Initiative (FMI).* U.S. Department of Transportation. "
        "https://www.bts.gov/explore-topics-and-geography/topics/freight-transportation/trucking-movements-bts-freight-mobility",
        "Bureau of Transportation Statistics. (2026). *BTS ATRI Freight Mobility Initiative: "
        "County-to-county truck travel times (2024)—Experimental* [Data set]. U.S. Department "
        "of Transportation. https://data.bts.gov/d/uta5-4eu5",
        "Kingma, D. P., & Ba, J. (2015). Adam: A method for stochastic optimization. "
        "*3rd International Conference on Learning Representations.* https://arxiv.org/abs/1412.6980",
        "Kirkpatrick, S., Gelatt, C. D., Jr., & Vecchi, M. P. (1983). Optimization by simulated "
        "annealing. *Science, 220*(4598), 671–680. https://doi.org/10.1126/science.220.4598.671",
        "Leslie, A., & Murray, D. (2024). *An analysis of the operational costs of trucking: "
        "2024 update.* American Transportation Research Institute. "
        "https://truckingresearch.org/2024/06/an-analysis-of-the-operational-costs-of-trucking-2024-update/",
        "SciPy. (n.d.-a). *linprog.* "
        "https://docs.scipy.org/doc/scipy-1.16.3/reference/generated/scipy.optimize.linprog.html",
        "SciPy. (n.d.-b). *milp.* "
        "https://docs.scipy.org/doc/scipy-1.16.3/reference/generated/scipy.optimize.milp.html",
        "U.S. Census Bureau. (n.d.). *Centers of population for the 2020 Census* [Data set]. "
        "https://www.census.gov/geographies/reference-files/2020/geo/2020-centers-population.html",
        "U.S. Census Bureau. (2023a). *Cartographic boundary files: States, 1:20,000,000* [Data set]. "
        "https://www2.census.gov/geo/tiger/GENZ2023/kml/cb_2023_us_state_20m.zip",
        "U.S. Census Bureau. (2023b). *Gazetteer files: 2023 counties* [Data set]. "
        "https://www.census.gov/geographies/reference-files/time-series/geo/gazetteer-files.2023.html",
        "U.S. Census Bureau. (2024a). *All sectors: County Business Patterns, including ZIP "
        "Code Business Patterns, by legal form of organization and employment size class "
        "for the U.S., states, and selected geographies: 2022* (Table CB2200CBP) [Data set]. "
        "https://data.census.gov/table/CBP2022.CB2200CBP",
        "U.S. Census Bureau. (2024b). *Total population: 2019–2023 American Community Survey "
        "5-year estimates* (Table B01003) [Data set]. https://data.census.gov/table/ACSDT5Y2023.B01003",
    ]
    for reference in references:
        R.reference(reference)

    R.h("Appendix: Reproducibility")
    R.p("Run python -m pytest -q for automated validation, python run_pipeline.py for the "
        "complete analysis from cached source files, and python report/build_docx.py for "
        "matching Word and Markdown reports. python verify_sources.py checks source values "
        "over the network and records hashes. --refresh-data downloads source files again; "
        "--quick uses fewer restarts and simulation scenarios and is not used for these results.")
    R.p(f"This run used Python {run['python']}, NumPy {run['numpy']}, and pandas {run['pandas']}. "
        f"Platform: {run['platform']}. The full analysis took {summary['runtime_total_s']:.1f} "
        "seconds. Seeds, algorithm settings, parameters, and source-verification state "
        "are stored in results/summary.json. Detailed decisions are in optimal_flows.csv, "
        "inventory_policy.csv, candidates.csv, and customers_demand.csv. The accompanying "
        "notebook runs the analysis and displays the same saved tables and figures.")
    return R.save()


if __name__ == "__main__":
    print("Wrote", build())
